import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime, date
import plotly.express as px
import plotly.graph_objects as go
from pathlib import Path
import shutil
import tempfile
from io import BytesIO
from PIL import Image
import base64
import html
import json
import hashlib
import secrets
import hmac

# =========================================================
# CONFIGURAÇÃO GERAL
# =========================================================
st.set_page_config(
    page_title="Dashboard PMO",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

CORES_STATUS = {
    "Concluida": "#16A34A",
    "Em Andamento": "#FACC15",
    "Em Atraso": "#DC2626",
    "Concluida atrasada": "#F97316",
    "Sem data alvo": "#6B7280",
}

STATUS_CURTO = {
    "Concluida": "No Prazo",
    "Em Andamento": "Em Atenção",
    "Em Atraso": "Atrasadas",
    "Concluida atrasada": "Concl. Atrasada",
    "Sem data alvo": "Sem Data",
}

# Ajuste aqui caso sua planilha mude de posição/colunas
ABA_PADRAO = "Cronogrma Macro"
LINHA_INICIAL = 15
LINHA_FINAL = 239
COL_ITEM = "C"
COL_ATIVIDADE = "D"
COL_RESPONSAVEL = "P"
COL_DATA_ALVO = "R"
COL_DATA_CONCLUSAO = "S"

# Logo fixo na Sidebar: mantenha o arquivo de imagem na mesma pasta deste programa.
LOGO_SIDEBAR = Path(__file__).with_name("AGSolution_logo.png")
VERSAO_DASHBOARD = "Rev0x_Demo"

# =========================================================
# CONTROLE DE ACESSO - LOGIN E USUÁRIOS
# =========================================================
USUARIOS_JSON = Path(__file__).with_name("usuarios.json")
USUARIO_ADMIN_PADRAO = "admin"
SENHA_ADMIN_PADRAO = "admin123"


def gerar_hash_senha(senha: str, salt=None) -> str:
    """Gera hash seguro da senha usando PBKDF2."""
    if salt is None:
        salt = secrets.token_hex(16)
    hash_senha = hashlib.pbkdf2_hmac(
        "sha256",
        senha.encode("utf-8"),
        salt.encode("utf-8"),
        120000,
    ).hex()
    return f"{salt}${hash_senha}"


def validar_senha(senha: str, senha_hash: str) -> bool:
    try:
        salt, hash_salvo = senha_hash.split("$", 1)
        hash_digitado = hashlib.pbkdf2_hmac(
            "sha256",
            senha.encode("utf-8"),
            salt.encode("utf-8"),
            120000,
        ).hex()
        return hmac.compare_digest(hash_digitado, hash_salvo)
    except Exception:
        return False


def carregar_usuarios() -> dict:
    """Carrega usuários do arquivo local. Cria admin padrão no primeiro uso."""
    if not USUARIOS_JSON.exists():
        usuarios_iniciais = {
            USUARIO_ADMIN_PADRAO: {
                "nome": "Administrador",
                "usuario": USUARIO_ADMIN_PADRAO,
                "senha_hash": gerar_hash_senha(SENHA_ADMIN_PADRAO),
                "perfil": "Administrador",
                "ativo": True,
                "criado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            }
        }
        salvar_usuarios(usuarios_iniciais)
        return usuarios_iniciais

    try:
        return json.loads(USUARIOS_JSON.read_text(encoding="utf-8"))
    except Exception:
        st.error("Erro ao ler usuarios.json. Verifique se o arquivo não está corrompido.")
        st.stop()


def salvar_usuarios(usuarios: dict) -> None:
    USUARIOS_JSON.write_text(
        json.dumps(usuarios, ensure_ascii=False, indent=4),
        encoding="utf-8",
    )


def fazer_logout() -> None:
    for chave in ["autenticado", "usuario_logado", "perfil_logado", "nome_logado"]:
        st.session_state.pop(chave, None)
    st.rerun()


def tela_login() -> None:
    st.markdown(
        """
        <style>
        .login-box {
            background: white;
            border: 1px solid #E5E7EB;
            border-radius: 18px;
            padding: 26px;
            box-shadow: 0 10px 30px rgba(15, 23, 42, 0.08);
            margin-top: 40px;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    col1, col2, col3 = st.columns([1, 1.15, 1])
    with col2:
        st.markdown("<div class='login-box'>", unsafe_allow_html=True)
        st.markdown("## 🔐 Acesso ao Dashboard PMO")
        st.caption("Informe usuário e senha para acessar o sistema.")

        with st.form("form_login"):
            usuario = st.text_input("Usuário")
            senha = st.text_input("Senha", type="password")
            entrar = st.form_submit_button("Entrar", width="stretch")

        if entrar:
            usuarios = carregar_usuarios()
            dados_usuario = usuarios.get(usuario.strip())
            if (
                dados_usuario
                and dados_usuario.get("ativo", True)
                and validar_senha(senha, dados_usuario.get("senha_hash", ""))
            ):
                st.session_state["autenticado"] = True
                st.session_state["usuario_logado"] = usuario.strip()
                st.session_state["perfil_logado"] = dados_usuario.get("perfil", "Usuário")
                st.session_state["nome_logado"] = dados_usuario.get("nome", usuario.strip())
                st.rerun()
            else:
                st.error("Usuário ou senha inválidos, ou usuário inativo.")

        st.info("Primeiro acesso: usuário admin / senha admin123. Altere a senha após publicar.")
        st.markdown("</div>", unsafe_allow_html=True)


def exigir_login() -> None:
    if not st.session_state.get("autenticado"):
        tela_login()
        st.stop()


def usuario_eh_admin() -> bool:
    return st.session_state.get("perfil_logado") == "Administrador"


def pagina_administrar_usuarios() -> None:
    st.markdown("<div class='section-title'>Administração de Usuários</div>", unsafe_allow_html=True)
    if not usuario_eh_admin():
        st.warning("Apenas usuário Administrador pode acessar esta área.")
        return

    usuarios = carregar_usuarios()

    st.subheader("Criar novo usuário")
    with st.form("form_novo_usuario"):
        c1, c2 = st.columns(2)
        with c1:
            novo_nome = st.text_input("Nome")
            novo_usuario = st.text_input("Usuário de login")
        with c2:
            novo_perfil = st.selectbox("Perfil", ["Usuário", "Administrador"])
            novo_ativo = st.checkbox("Usuário ativo", value=True)
        nova_senha = st.text_input("Senha inicial", type="password")
        criar = st.form_submit_button("Criar usuário", width="stretch")

    if criar:
        novo_usuario = novo_usuario.strip()
        if not novo_nome.strip() or not novo_usuario or not nova_senha:
            st.error("Preencha nome, usuário e senha.")
        elif novo_usuario in usuarios:
            st.error("Este usuário já existe.")
        else:
            usuarios[novo_usuario] = {
                "nome": novo_nome.strip(),
                "usuario": novo_usuario,
                "senha_hash": gerar_hash_senha(nova_senha),
                "perfil": novo_perfil,
                "ativo": novo_ativo,
                "criado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            }
            salvar_usuarios(usuarios)
            st.success("Usuário criado com sucesso.")
            st.rerun()

    st.divider()
    st.subheader("Usuários cadastrados")
    tabela = []
    for u, d in usuarios.items():
        tabela.append({
            "Usuário": u,
            "Nome": d.get("nome", ""),
            "Perfil": d.get("perfil", "Usuário"),
            "Ativo": "Sim" if d.get("ativo", True) else "Não",
            "Criado em": d.get("criado_em", ""),
        })
    st.dataframe(pd.DataFrame(tabela), width="stretch", hide_index=True)

    st.divider()
    st.subheader("Alterar usuário existente")
    usuario_sel = st.selectbox("Selecione o usuário", sorted(usuarios.keys()))
    dados = usuarios[usuario_sel]

    with st.form("form_editar_usuario"):
        edit_nome = st.text_input("Nome", value=dados.get("nome", ""))
        edit_perfil = st.selectbox(
            "Perfil",
            ["Usuário", "Administrador"],
            index=0 if dados.get("perfil", "Usuário") == "Usuário" else 1,
        )
        edit_ativo = st.checkbox("Ativo", value=dados.get("ativo", True))
        edit_senha = st.text_input("Nova senha (deixe vazio para manter)", type="password")
        col_a, col_b = st.columns(2)
        with col_a:
            salvar = st.form_submit_button("Salvar alterações", width="stretch")
        with col_b:
            excluir = st.form_submit_button("Excluir usuário", width="stretch")

    if salvar:
        usuarios[usuario_sel]["nome"] = edit_nome.strip()
        usuarios[usuario_sel]["perfil"] = edit_perfil
        usuarios[usuario_sel]["ativo"] = edit_ativo
        if edit_senha:
            usuarios[usuario_sel]["senha_hash"] = gerar_hash_senha(edit_senha)
        salvar_usuarios(usuarios)
        st.success("Usuário atualizado com sucesso.")
        st.rerun()

    if excluir:
        if usuario_sel == st.session_state.get("usuario_logado"):
            st.error("Você não pode excluir o próprio usuário logado.")
        elif usuario_sel == USUARIO_ADMIN_PADRAO:
            st.error("O usuário admin padrão não pode ser excluído. Você pode alterar a senha ou desativar depois de criar outro administrador.")
        else:
            usuarios.pop(usuario_sel, None)
            salvar_usuarios(usuarios)
            st.success("Usuário excluído com sucesso.")
            st.rerun()


# =========================================================
# CSS - VISUAL ESTILO DASHBOARD WEB
# =========================================================
st.markdown(
    """
    <style>
    .main {background-color: #F5F7FB;}
    section[data-testid="stSidebar"] {background: linear-gradient(180deg, #061B3A 0%, #082B57 100%);}
    section[data-testid="stSidebar"] * {color: white !important;}
    section[data-testid="stSidebar"] .block-container {padding-top: 0.45rem; padding-left: 1rem; padding-right: 1rem;}
    .block-container {padding-top: 2.6rem; padding-bottom: 1rem;}
    .card {
        background: white;
        border: 1px solid #E5E7EB;
        border-radius: 16px;
        padding: 18px 20px;
        box-shadow: 0 6px 20px rgba(15, 23, 42, 0.05);
        min-height: 112px;
    }
    .metric-label {font-size: 13px; color: #475569; font-weight: 600;}
    .metric-value {font-size: 31px; color: #0F172A; font-weight: 800; line-height: 1.1;}
    .metric-delta {font-size: 13px; font-weight: 700; margin-top: 4px;}
    .dashboard-header {
        background: transparent;
        padding: 8px 0 12px 0;
        margin-top: 6px;
        margin-bottom: 8px;
        overflow: visible;
    }
    .header-title {
        font-size: 34px;
        font-weight: 850;
        color: #0F172A;
        margin: 0;
        padding: 4px 0 2px 0;
        line-height: 1.25;
        overflow: visible;
    }
    .header-subtitle {
        font-size: 16px;
        color: #475569;
        margin: 0;
        padding-top: 2px;
        line-height: 1.35;
    }
    .logo-box {
        display: flex;
        align-items: center;
        justify-content: flex-start;
        min-height: 74px;
        padding-top: 2px;
    }
    .client-logo-space {
        min-height: 96px;
        display: flex;
        align-items: center;
        justify-content: flex-start;
        padding-top: 4px;
    }

    .sidebar-logo-box {
        width: 100%;
        display: flex;
        justify-content: flex-start;
        align-items: center;
        padding: 0 0 8px 0;
        margin: 0 0 8px 0;
    }
    .sidebar-caption {
        font-size: 12px;
        color: #D8E6FF;
        line-height: 1.25;
        margin-top: -4px;
        margin-bottom: 14px;
    }
    .executive-header {
        background: #FFFFFF;
        border: 1px solid #E5E7EB;
        border-radius: 18px;
        box-shadow: 0 6px 20px rgba(15, 23, 42, 0.06);
        padding: 18px 22px;
        margin: 2px 0 18px 0;
        display: flex;
        align-items: center;
        gap: 20px;
        min-height: 118px;
        overflow: visible;
    }
    .executive-logo-area {
        width: 145px;
        min-width: 145px;
        display: flex;
        align-items: center;
        justify-content: center;
        border-right: 1px solid #E5E7EB;
        padding-right: 16px;
        min-height: 86px;
    }
    .executive-logo-area img {
        max-width: 125px;
        max-height: 78px;
        object-fit: contain;
    }
    .executive-logo-placeholder {
        width: 118px;
        height: 68px;
        border: 1px dashed #CBD5E1;
        border-radius: 12px;
        color: #94A3B8;
        font-size: 12px;
        display: flex;
        align-items: center;
        justify-content: center;
        text-align: center;
        padding: 6px;
    }
    .executive-main-area {
        flex: 1;
        min-width: 0;
    }
    .executive-title-row {
        display: flex;
        justify-content: space-between;
        align-items: flex-start;
        gap: 16px;
    }
    .executive-title {
        font-size: 36px;
        font-weight: 850;
        color: #0F172A;
        line-height: 1.18;
        margin: 0;
        padding: 0;
    }
    .executive-subtitle {
        font-size: 15px;
        color: #475569;
        margin-top: 2px;
        line-height: 1.35;
    }
    .executive-version {
        background: #EFF6FF;
        color: #1D4ED8;
        border: 1px solid #BFDBFE;
        border-radius: 999px;
        padding: 6px 14px;
        font-size: 14px;
        font-weight: 800;
        white-space: nowrap;
    }
    .executive-meta {
        display: grid;
        grid-template-columns: 1fr 1fr;
        gap: 6px 22px;
        margin-top: 10px;
        color: #475569;
        font-size: 13.5px;
        line-height: 1.35;
    }
    .executive-meta b {color: #0F172A;}
    @media (max-width: 900px) {
        .executive-header {align-items: flex-start; gap: 12px; padding: 14px;}
        .executive-logo-area {width: 105px; min-width: 105px; padding-right: 10px;}
        .executive-title {font-size: 28px;}
        .executive-meta {grid-template-columns: 1fr;}
    }
    .section-title {font-size: 19px; font-weight: 800; color:#0F172A; margin: 0 0 10px 0;}
    div[data-testid="stMetric"] {background: white; border-radius: 14px; padding: 14px; border: 1px solid #E5E7EB;}


    /* =======================================================
       SIDEBAR CORPORATIVA PMO - CAMPOS, UPLOADS E BOTÕES
    ======================================================= */
    section[data-testid="stSidebar"] {
        background: linear-gradient(180deg,#061B3A 0%,#082B57 100%) !important;
    }

    section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] p,
    section[data-testid="stSidebar"] span,
    section[data-testid="stSidebar"] div {
        color: #FFFFFF !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stTextInput"] input {
        background-color: #0F3D73 !important;
        color: #FFFFFF !important;
        border: 1px solid #4A90E2 !important;
        border-radius: 10px !important;
        font-weight: 500 !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stTextInput"] input::placeholder {
        color: #D0E3FF !important;
    }

    section[data-testid="stSidebar"] [data-testid="stFileUploader"] {
        background-color: #0F3D73 !important;
        border: 1px solid #4A90E2 !important;
        border-radius: 12px !important;
        padding: 12px !important;
    }

    section[data-testid="stSidebar"] [data-testid="stFileUploader"] section {
        background-color: #0F3D73 !important;
        border: 1px dashed #4A90E2 !important;
        border-radius: 10px !important;
    }

    section[data-testid="stSidebar"] [data-testid="stFileUploader"] * {
        color: #FFFFFF !important;
    }

    section[data-testid="stSidebar"] button {
        background-color: #2563EB !important;
        color: #FFFFFF !important;
        border: none !important;
        border-radius: 10px !important;
        font-weight: 700 !important;
        transition: 0.3s ease-in-out !important;
    }

    section[data-testid="stSidebar"] button:hover {
        background-color: #1D4ED8 !important;
        color: #FFFFFF !important;
    }

    section[data-testid="stSidebar"] button[kind="secondary"] {
        background-color: #2563EB !important;
        color: #FFFFFF !important;
    }

    section[data-testid="stSidebar"] [role="radiogroup"] {
        background-color: #0F3D73 !important;
        border: 1px solid #4A90E2 !important;
        border-radius: 12px !important;
        padding: 10px !important;
    }

    section[data-testid="stSidebar"] hr {
        border-color: rgba(255,255,255,0.15) !important;
    }

    .sidebar-caption {
        color: #D0E3FF !important;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# Exige autenticação antes de carregar o dashboard
exigir_login()

# =========================================================
# FUNÇÕES
# =========================================================
def carregar_workbook(arquivo_origem):
    try:
        temp_dir = Path(tempfile.gettempdir())
        arquivo_temp = temp_dir / f"pmo_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        if hasattr(arquivo_origem, "read"):
            arquivo_temp.write_bytes(arquivo_origem.getvalue())
        else:
            shutil.copy2(Path(arquivo_origem), arquivo_temp)

        return load_workbook(arquivo_temp, data_only=True)

    except PermissionError:
        st.error("Feche o arquivo Excel antes de executar. Ele está bloqueado pelo Excel ou OneDrive.")
        st.stop()
    except FileNotFoundError:
        st.error(f"Arquivo não encontrado: {arquivo_origem}")
        st.stop()
    except Exception as erro:
        st.error(f"Erro ao carregar a planilha: {erro}")
        st.stop()


def normalizar_data(valor):
    return pd.to_datetime(valor, dayfirst=True, errors="coerce")


def calcular_status(row):
    data_alvo = row["Data Alvo"]
    data_atividade = row["Data Atividade"]
    hoje = pd.Timestamp(datetime.now().date())

    if pd.notna(data_atividade) and pd.notna(data_alvo) and data_atividade > data_alvo:
        return "Concluida atrasada"
    if pd.notna(data_atividade) and pd.notna(data_alvo) and data_atividade <= data_alvo:
        return "Concluida"
    if pd.isna(data_atividade) and pd.notna(data_alvo) and data_alvo < hoje:
        return "Em Atraso"
    if pd.isna(data_atividade) and pd.notna(data_alvo) and data_alvo >= hoje:
        return "Em Andamento"
    return "Sem data alvo"


def extrair_dados(wb, aba):
    if aba not in wb.sheetnames:
        st.error(f"Aba '{aba}' não encontrada. Abas disponíveis: {', '.join(wb.sheetnames)}")
        st.stop()

    ws = wb[aba]
    dados = []

    for linha in range(LINHA_INICIAL, LINHA_FINAL):
        item = ws[f"{COL_ITEM}{linha}"].value
        atividade = ws[f"{COL_ATIVIDADE}{linha}"].value

        if atividade is None:
            continue

        alvo = ws[f"{COL_DATA_ALVO}{linha}"].value
        data_atividade = ws[f"{COL_DATA_CONCLUSAO}{linha}"].value

        responsavel = ws[f"{COL_RESPONSAVEL}{linha}"].value
        fase = ws[f"B{linha}"].value if ws.max_column >= 2 else None

        dados.append({
            "Linha": linha,
            "Item": item,
            "Fase": fase if fase else "Não informado",
            "Atividade": atividade,
            "Responsável": responsavel if responsavel else "Não informado",
            "Data Alvo": alvo,
            "Data Atividade": data_atividade,
        })

    df = pd.DataFrame(dados)
    if df.empty:
        st.warning("Nenhuma atividade encontrada na faixa configurada.")
        st.stop()

    # Normaliza colunas de texto para evitar erro: TypeError '<' not supported between float and str
    for col in ["Fase", "Responsável", "Atividade", "Item"]:
        if col in df.columns:
            df[col] = df[col].apply(lambda x: limpar_texto(x, padrao="Não informado"))

    # Conversão robusta das datas: evita erro ao subtrair numpy.ndarray/object de Timestamp.
    # Trata datas do Excel, textos, células vazias e valores inválidos como NaT.
    df["Data Alvo"] = pd.to_datetime(df["Data Alvo"], errors="coerce", dayfirst=True)
    df["Data Atividade"] = pd.to_datetime(df["Data Atividade"], errors="coerce", dayfirst=True)

    df["Status"] = df.apply(calcular_status, axis=1)

    hoje = pd.Timestamp.now().normalize()
    df["Dias para Vencer"] = pd.NA
    mask_data_alvo = df["Data Alvo"].notna()
    df.loc[mask_data_alvo, "Dias para Vencer"] = (
        df.loc[mask_data_alvo, "Data Alvo"] - hoje
    ).dt.days.astype("Int64")

    df["Mês Alvo"] = df["Data Alvo"].dt.to_period("M").astype(str).replace("NaT", "Sem data")
    df["Semana Alvo"] = df["Data Alvo"].dt.strftime("%d/%m/%Y").fillna("Sem data")
    df["Inicio Gantt"] = df["Data Alvo"] - pd.Timedelta(days=7)
    df["Fim Gantt"] = df["Data Atividade"].fillna(df["Data Alvo"])
    df.loc[df["Fim Gantt"] < df["Inicio Gantt"], "Fim Gantt"] = df["Inicio Gantt"] + pd.Timedelta(days=1)
    return df


def pct(parte, total):
    return (parte / total * 100) if total else 0


def card(titulo, valor, subtitulo, cor="#0F172A", icone=""):
    st.markdown(
        f"""
        <div class="card">
            <div class="metric-label">{icone} {titulo}</div>
            <div class="metric-value">{valor}</div>
            <div class="metric-delta" style="color:{cor};">{subtitulo}</div>
        </div>
        """,
        unsafe_allow_html=True
    )


def gerar_excel(df, resumo):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Atividades")
        resumo.to_excel(writer, index=False, sheet_name="Resumo")
    return output.getvalue()


def limpar_texto(valor, padrao="Não informado"):
    """Converte valores vindos do Excel para texto seguro, evitando mistura float/str nos filtros."""
    if pd.isna(valor):
        return padrao
    texto = str(valor).strip()
    if texto == "" or texto.lower() in ["nan", "nat", "none"]:
        return padrao
    # remove .0 de números inteiros vindos do Excel, ex.: 10.0 -> 10
    try:
        numero = float(texto.replace(",", "."))
        if numero.is_integer():
            return str(int(numero))
    except Exception:
        pass
    return texto


def lista_filtro(df, coluna):
    """Lista segura para multiselect: remove nulos, converte tudo para string e ordena."""
    if coluna not in df.columns:
        return []
    valores = [limpar_texto(v) for v in df[coluna].tolist()]
    return sorted(set(v for v in valores if v))


def arquivo_upload_para_base64(arquivo):
    """Converte imagem carregada no Streamlit para base64, permitindo posicionamento no cabeçalho HTML."""
    if arquivo is None:
        return None
    try:
        return base64.b64encode(arquivo.getvalue()).decode("utf-8")
    except Exception:
        return None

# =========================================================
# SIDEBAR
# =========================================================
with st.sidebar:
    st.markdown("<div class='sidebar-logo-box'>", unsafe_allow_html=True)
    if LOGO_SIDEBAR.exists():
        st.image(str(LOGO_SIDEBAR), width=230)
    else:
        st.markdown("## AGSolution")
        st.caption("Arquivo AGSolution_logo.png não encontrado na pasta do programa.")
    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown(
        """
        <div class='sidebar-caption'>
            
        </div>
        """,
        unsafe_allow_html=True
    )

    st.divider()

    opcoes_menu = ["Visão Geral", "Cronograma", "Projetos", "Indicadores", "Riscos", "Relatórios", "Exportar", "Configurações"]
    if usuario_eh_admin():
        opcoes_menu.append("Administração de Usuários")

    menu = st.radio(
        "Menu",
        opcoes_menu,
        index=0
    )

    st.caption(f"Logado como: {st.session_state.get('nome_logado', '')} | {st.session_state.get('perfil_logado', '')}")
    if st.button("Sair", width="stretch"):
        fazer_logout()

    st.divider()
    arquivo_upload = st.file_uploader("Carregar planilha Excel", type=["xlsx"])

    st.divider()
    cliente_nome = st.text_input(
        "Nome do Cliente",
        value="         ",
        key="cliente_nome"
    )

    logo_cliente = st.file_uploader(
        "Carregar logo do cliente",
        type=["png", "jpg", "jpeg"],
        key="logo_cliente"
    )

# =========================================================
# CARREGAMENTO
# =========================================================
if arquivo_upload is None:
    st.warning("Carregue a planilha Excel para iniciar o Dashboard PMO.")
    st.stop()

wb = carregar_workbook(arquivo_upload)
aba_selecionada = ABA_PADRAO if ABA_PADRAO in wb.sheetnames else wb.sheetnames[0]
df_base = extrair_dados(wb, aba_selecionada)

# =========================================================
# CABEÇALHO EXECUTIVO E FILTROS GLOBAIS
# =========================================================
nome_arquivo_exibicao = arquivo_upload.name
ultima_atualizacao = datetime.now().strftime("%d/%m/%Y %H:%M")
cliente_nome_seguro = html.escape(cliente_nome.strip() if cliente_nome else "Cliente não informado")
nome_arquivo_seguro = html.escape(nome_arquivo_exibicao)
logo_cliente_b64 = arquivo_upload_para_base64(logo_cliente)

if logo_cliente_b64:
    logo_cliente_html = f"<img src='data:image/png;base64,{logo_cliente_b64}' alt='Logo do cliente'>"
else:
    logo_cliente_html = "<div class='executive-logo-placeholder'>Logo do cliente</div>"

st.markdown(
    f"""
    <div class='executive-header'>
        <div class='executive-logo-area'>
            {logo_cliente_html}
        </div>
        <div class='executive-main-area'>
            <div class='executive-title-row'>
                <div>
                    <div class='executive-title'>Dashboard PMO</div>
                    <div class='executive-subtitle'>Visão geral do desempenho dos projetos e cronograma</div>
                </div>
                <div class='executive-version'>{VERSAO_DASHBOARD}</div>
            </div>
            <div class='executive-meta'>
                <div><b>Cliente:</b> {cliente_nome_seguro}</div>
                <div><b>Última atualização:</b> {ultima_atualizacao}</div>
                <div style='grid-column: 1 / -1;'><b>Arquivo:</b> {nome_arquivo_seguro}</div>
            </div>
        </div>
    </div>
    """,
    unsafe_allow_html=True
)

f1, f2, f3, f4 = st.columns([1.2, 1.2, 1.2, 1.8])
with f1:
    status_opcoes = lista_filtro(df_base, "Status")
    status_sel = st.multiselect("Status", options=status_opcoes, default=status_opcoes)
with f2:
    resp_opcoes = lista_filtro(df_base, "Responsável")
    resp_sel = st.multiselect("Responsável", options=resp_opcoes, default=resp_opcoes)
with f3:
    fase_opcoes = lista_filtro(df_base, "Fase")
    fase_sel = st.multiselect("Fase", options=fase_opcoes, default=fase_opcoes)
with f4:
    datas_validas = df_base["Data Alvo"].dropna()
    if not datas_validas.empty:
        periodo = st.date_input(
            "Período Data Alvo",
            value=(datas_validas.min().date(), datas_validas.max().date()),
            format="DD/MM/YYYY"
        )
    else:
        periodo = None

df = df_base[
    df_base["Status"].isin(status_sel)
    & df_base["Responsável"].isin(resp_sel)
    & df_base["Fase"].isin(fase_sel)
].copy()

if periodo and len(periodo) == 2:
    inicio, fim = pd.Timestamp(periodo[0]), pd.Timestamp(periodo[1])
    df = df[(df["Data Alvo"].isna()) | ((df["Data Alvo"] >= inicio) & (df["Data Alvo"] <= fim))]

# =========================================================
# MÉTRICAS
# =========================================================
total = len(df)
qtd_concluida = int((df["Status"] == "Concluida").sum())
qtd_andamento = int((df["Status"] == "Em Andamento").sum())
qtd_atraso = int((df["Status"] == "Em Atraso").sum())
qtd_concluida_atrasada = int((df["Status"] == "Concluida atrasada").sum())
qtd_atencao = qtd_andamento + qtd_concluida_atrasada
conclusao_geral = pct(qtd_concluida + qtd_concluida_atrasada, total)

resumo = df.groupby("Status", dropna=False).size().reset_index(name="Quantidade")
resumo["Percentual"] = (resumo["Quantidade"] / total * 100).round(1) if total else 0
resumo["Status Curto"] = resumo["Status"].map(STATUS_CURTO).fillna(resumo["Status"])

m1, m2, m3, m4, m5 = st.columns(5)
with m1: card("Total de Atividades", total, "100% do total", "#2563EB", "📋")
with m2: card("Atrasadas", qtd_atraso, f"{pct(qtd_atraso, total):.1f}% do total", "#DC2626", "⏰")
with m3: card("Em Atenção", qtd_atencao, f"{pct(qtd_atencao, total):.1f}% do total", "#D97706", "⚠️")
with m4: card("No Prazo", qtd_concluida, f"{pct(qtd_concluida, total):.1f}% do total", "#16A34A", "✅")
with m5: card("Conclusão Geral", f"{conclusao_geral:.0f}%", "% de atividades concluídas", "#2563EB", "🎯")

# =========================================================
# PÁGINAS
# =========================================================
if menu == "Visão Geral":
    c1, c2 = st.columns([1, 1.35])
    with c1:
        st.markdown("<div class='section-title'>Status das Atividades</div>", unsafe_allow_html=True)
        fig_pizza = px.pie(
            resumo,
            values="Quantidade",
            names="Status Curto",
            hole=0.52,
            color="Status",
            color_discrete_map=CORES_STATUS,
        )
        fig_pizza.update_traces(textinfo="percent", textfont_size=15)
        fig_pizza.update_layout(height=390, margin=dict(l=10, r=10, t=20, b=10), showlegend=True)
        st.plotly_chart(fig_pizza, width="stretch")

    with c2:
        st.markdown("<div class='section-title'>Evolução de Conclusão (%)</div>", unsafe_allow_html=True)
        evolucao = df.dropna(subset=["Data Alvo"]).sort_values("Data Alvo").copy()
        evolucao["Concluida"] = evolucao["Status"].isin(["Concluida", "Concluida atrasada"]).astype(int)
        if not evolucao.empty:
            evolucao["Acumulado"] = evolucao["Concluida"].cumsum() / pd.Series(range(1, len(evolucao) + 1), index=evolucao.index)
            evolucao["Conclusão %"] = (evolucao["Acumulado"] * 100).round(1)
            fig_linha = px.line(evolucao, x="Data Alvo", y="Conclusão %", markers=True)
            fig_linha.update_layout(height=390, yaxis_range=[0, 100], margin=dict(l=10, r=10, t=20, b=10))
            st.plotly_chart(fig_linha, width="stretch")
        else:
            st.info("Sem datas alvo para gerar evolução.")

    c3, c4, c5 = st.columns([1.45, 0.7, 0.85])
    with c3:
        st.markdown("<div class='section-title'>Cronograma Macro (Gantt)</div>", unsafe_allow_html=True)
        gantt = df.dropna(subset=["Inicio Gantt", "Fim Gantt"]).head(35)
        if not gantt.empty:
            fig_gantt = px.timeline(
                gantt,
                x_start="Inicio Gantt",
                x_end="Fim Gantt",
                y="Atividade",
                color="Status",
                color_discrete_map=CORES_STATUS,
                hover_data=["Item", "Responsável", "Data Alvo", "Data Atividade"],
            )
            fig_gantt.update_yaxes(autorange="reversed")
            fig_gantt.update_layout(height=380, margin=dict(l=10, r=10, t=20, b=10), showlegend=False)
            st.plotly_chart(fig_gantt, width="stretch")
        else:
            st.info("Sem dados para Gantt.")

    with c4:
        st.markdown("<div class='section-title'>Atividades por Responsável</div>", unsafe_allow_html=True)
        por_resp = df.groupby("Responsável").size().reset_index(name="Quantidade").sort_values("Quantidade")
        fig_resp = px.bar(por_resp, x="Quantidade", y="Responsável", orientation="h", text="Quantidade")
        fig_resp.update_layout(height=380, margin=dict(l=10, r=10, t=20, b=10))
        st.plotly_chart(fig_resp, width="stretch")

    with c5:
        st.markdown("<div class='section-title'>Principais Riscos</div>", unsafe_allow_html=True)
        riscos = df[df["Status"].isin(["Em Atraso", "Concluida atrasada"])]\
            .sort_values(["Status", "Dias para Vencer"]).head(5)
        if riscos.empty:
            st.success("Nenhum risco crítico no filtro atual.")
        else:
            for _, r in riscos.iterrows():
                st.markdown(
                    f"**🔴 {r['Atividade']}**  \n"
                    f"Status: {r['Status']}  \n"
                    f"Responsável: {r['Responsável']}  \n"
                    f"Alvo: {r['Data Alvo'].strftime('%d/%m/%Y') if pd.notna(r['Data Alvo']) else 'Sem data'}"
                )
                st.divider()

elif menu == "Cronograma":
    st.markdown("<div class='section-title'>Cronograma Detalhado</div>", unsafe_allow_html=True)
    gantt = df.dropna(subset=["Inicio Gantt", "Fim Gantt"])
    if gantt.empty:
        st.info("Sem dados para exibir no Gantt com o filtro atual.")
    else:
        fig_gantt = px.timeline(
            gantt,
            x_start="Inicio Gantt",
            x_end="Fim Gantt",
            y="Atividade",
            color="Status",
            color_discrete_map=CORES_STATUS,
            hover_data=["Item", "Responsável", "Data Alvo", "Data Atividade", "Dias para Vencer"],
        )
        fig_gantt.update_yaxes(autorange="reversed")
        fig_gantt.update_layout(height=max(500, len(gantt) * 18), margin=dict(l=10, r=10, t=20, b=10))
        st.plotly_chart(fig_gantt, width="stretch")

elif menu == "Projetos":
    st.markdown("<div class='section-title'>Visão por Fase / Projeto</div>", unsafe_allow_html=True)
    matriz = pd.crosstab(df["Fase"], df["Status"]).reset_index()
    st.dataframe(matriz, width="stretch", hide_index=True)
    fig_fase = px.bar(df, x="Fase", color="Status", color_discrete_map=CORES_STATUS, barmode="stack")
    st.plotly_chart(fig_fase, width="stretch")

elif menu == "Indicadores":
    st.markdown("<div class='section-title'>Indicadores Avançados</div>", unsafe_allow_html=True)
    ind1, ind2 = st.columns(2)
    with ind1:
        fig_mes = px.histogram(df.dropna(subset=["Data Alvo"]), x="Mês Alvo", color="Status", color_discrete_map=CORES_STATUS, barmode="group")
        fig_mes.update_layout(title="Atividades por Mês Alvo")
        st.plotly_chart(fig_mes, width="stretch")
    with ind2:
        heat = pd.crosstab(df["Responsável"], df["Status"])
        fig_heat = px.imshow(heat, text_auto=True, aspect="auto", title="Mapa de Calor por Responsável x Status")
        st.plotly_chart(fig_heat, width="stretch")

elif menu == "Riscos":
    st.markdown("<div class='section-title'>Riscos e Pendências</div>", unsafe_allow_html=True)
    riscos = df[df["Status"].isin(["Em Atraso", "Concluida atrasada", "Sem data alvo"])].copy()
    riscos["Criticidade"] = riscos["Status"].map({"Em Atraso": "Alto", "Concluida atrasada": "Médio", "Sem data alvo": "Médio"}).fillna("Baixo")
    st.dataframe(riscos[["Item", "Atividade", "Responsável", "Data Alvo", "Data Atividade", "Status", "Criticidade"]], width="stretch", hide_index=True)

elif menu == "Relatórios":
    st.markdown("<div class='section-title'>Relatório Executivo</div>", unsafe_allow_html=True)
    st.write(f"Total de atividades avaliadas: **{total}**")
    st.write(f"Conclusão geral: **{conclusao_geral:.1f}%**")
    st.write(f"Atividades em atraso: **{qtd_atraso}**")
    st.write(f"Atividades em atenção: **{qtd_atencao}**")
    st.dataframe(resumo, width="stretch", hide_index=True)

elif menu == "Exportar":
    st.markdown("<div class='section-title'>Exportação</div>", unsafe_allow_html=True)
    excel_bytes = gerar_excel(df, resumo)
    st.download_button(
        "⬇️ Baixar relatório Excel",
        data=excel_bytes,
        file_name=f"Relatorio_PMO_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    csv = df.to_csv(index=False, sep=";").encode("utf-8-sig")
    st.download_button("⬇️ Baixar atividades CSV", data=csv, file_name="atividades_pmo.csv", mime="text/csv")

elif menu == "Configurações":
    st.markdown("<div class='section-title'>Configurações do Sistema</div>", unsafe_allow_html=True)
    st.info("Para mudar colunas, linhas ou aba padrão, ajuste as constantes no início do arquivo Python.")
    st.code(
        f"ABA_PADRAO = '{ABA_PADRAO}'\nLINHA_INICIAL = {LINHA_INICIAL}\nLINHA_FINAL = {LINHA_FINAL}\nCOL_RESPONSAVEL = '{COL_RESPONSAVEL}'\nCOL_DATA_ALVO = '{COL_DATA_ALVO}'\nCOL_DATA_CONCLUSAO = '{COL_DATA_CONCLUSAO}'"
    )

elif menu == "Administração de Usuários":
    pagina_administrar_usuarios()

# =========================================================
# TABELA FINAL COM CORES FIXAS
# =========================================================
if menu != "Administração de Usuários":
    st.divider()
    st.markdown("<div class='section-title'>Atividades</div>", unsafe_allow_html=True)

    colunas_exibir = ["Item", "Fase", "Atividade", "Responsável", "Data Alvo", "Data Atividade", "Dias para Vencer", "Status"]

    for col in ["Data Alvo", "Data Atividade"]:
        df[col] = pd.to_datetime(df[col], errors="coerce")

    def colorir_status(row):
        cor = CORES_STATUS.get(row["Status"], "#FFFFFF")
        texto = "#FFFFFF" if row["Status"] in ["Em Atraso", "Concluida atrasada", "Sem data alvo"] else "#111827"
        return [f"background-color: {cor}; color: {texto}; font-weight: 600" if c == "Status" else "" for c in row.index]

    st.dataframe(
        df[colunas_exibir].style.apply(colorir_status, axis=1).format({
            "Data Alvo": lambda x: x.strftime("%d/%m/%Y") if pd.notna(x) else "",
            "Data Atividade": lambda x: x.strftime("%d/%m/%Y") if pd.notna(x) else "",
        }),
        width="stretch",
        hide_index=True
    )

    st.caption(f"{VERSAO_DASHBOARD} | Última atualização: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Cliente: {cliente_nome if cliente_nome else 'Cliente não informado'} | Fonte: {arquivo_upload.name}")
